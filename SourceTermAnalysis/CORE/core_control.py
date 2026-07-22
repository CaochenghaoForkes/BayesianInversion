import os
import shutil
import numpy as np
import pandas as pd
import bisect
from pathlib import Path
from typing import List,TypeVar
from CORE.core_xml_reader import CoreXMLReader
from CORE.core_tracer import Tracer
from openpyxl import Workbook
from TOOL.cosmos_general_function import round_sig,formatting_line
from PIL import Image
import matplotlib.pyplot as plt 

T = TypeVar("T")

class CoreControl:
    '''
        控制程序的输出
    '''
    
    class Plot:
        '''
            对输出结果绘图
        '''
        @classmethod
        def init_plot(cls,output_dir_raw,accident_begin,x_list,y_list):
            # 基本设置
            cls.output_dir = os.path.join(output_dir_raw,'Plot')
            if os.path.exists(cls.output_dir):
                shutil.rmtree(cls.output_dir)
            os.mkdir(cls.output_dir)
            cls.accident_begin = accident_begin
            cls.x_list = np.asarray(x_list)*100
            cls.y_list = np.asarray(y_list)*100
            
            # 绘图风格
            cls.x_2d = 6
            cls.y_2d = 12
            cls.x_1d = 6
            cls.y_1d = 4
            cls.dpi_1d = 300
            cls.dpi_2d = 300
            cls.linestyle='-'
            cls.linecolor='black'
            cls.linewidth=2.0
            cls.marker='v'
            cls.markersize=6
            cls.markerfacecolor='black'
            cls.markeredgewidth=1.5
            cls.markeredgecolor='#8B0000'
            cls.point_number = 50
            cls.axvline_color = 'blue'
            cls.axvline_style = '--'
            cls.axvline_width = 1.0
            cls.cmap = 'viridis'
            cls.duration = 50
            plt.rcParams['font.size'] = 14
            plt.rcParams['font.family'] = 'Times New Roman'
        
        @classmethod
        def plot_line(cls,y_label,y_unit,y,x,save_name):
            '''
                绘制曲线
            '''
            plt.figure(figsize=(cls.x_1d,cls.y_1d), dpi=cls.dpi_1d)
            point_number = cls.point_number
            interval = int(len(x) / point_number) if len(x)>2*point_number else 1
            plt.plot(x, y,
                color=cls.linecolor,
                linewidth=cls.linewidth,
                linestyle=cls.linestyle,
                marker=cls.marker,
                markersize=cls.markersize,
                markerfacecolor=cls.markerfacecolor,
                markeredgewidth=cls.markeredgewidth,
                markeredgecolor=cls.markeredgecolor,
                markevery=interval)
            if x[-1] > cls.accident_begin:
                plt.axvline(cls.accident_begin,color=cls.axvline_color,linestyle=cls.axvline_style,linewidth=cls.axvline_width)
            plt.ylabel(y_label+y_unit)
            plt.xlabel('Time([s])')
            plt.xlim(cls.accident_begin-(x[-1]-cls.accident_begin)*0.1,x[-1]+(x[-1]-cls.accident_begin)*0.1)
            plt.tight_layout()
            plt.savefig(os.path.join(cls.output_dir,f'{save_name}'+'.png'))
            plt.close()

        @classmethod
        def plot_2d(cls,matrix,z_label,z_unit,z_max,z_min,target_time,save_name):
            '''
                绘制二维图
            '''
            plt.figure(figsize=(cls.x_2d,cls.y_2d), dpi=cls.dpi_2d)
            pc = plt.pcolormesh(cls.x_list,cls.y_list,matrix,shading='auto',cmap=cls.cmap,vmin=z_min,vmax=z_max)  
            plt.gca().invert_yaxis()
            cbar = plt.colorbar(pc)
            cbar.set_label(z_label+z_unit)
            plt.xlabel('Radius([cm])')
            plt.ylabel('Depth([cm])')
            if target_time == 'steady':
                title = 'Steady'
            else:
                title = f'Time:{target_time}([s])'
            plt.title(title)
            plt.tight_layout()
            plt.savefig(os.path.join(cls.output_dir,f'{save_name}_{target_time}'+'.png'))
            plt.close()
        
        @classmethod
        def _png_to_gif(cls,files,name):
            '''
                png转gif
            '''
            images = []
            for file_path in files:
                img = Image.open(file_path)
                images.append(img)
            images[0].save(os.path.join(cls.output_dir,name+'.gif'),save_all=True,append_images=images[1:],duration=cls.duration,loop=0,optimize=True)
            
    def __init__(self,reader:CoreXMLReader,main_control=True):
        # 基本信息
        self.output_dir_raw = reader.control_dict['output_path']
        self.output_type = reader.control_dict['output_type']
        self.accident_begin = reader.control_dict['accident_begin']
        self.element_number = reader.core_properties_dict['element_number'] 
        self.output_nuclide = reader.control_dict['output_nuclide']
        self.time_list = reader.control_dict['time']
        self.time_step_list = reader.control_dict['time_step']
        self.write_interval_list = reader.control_dict['write_interval']
        self.mars_input_path = reader.core_mars_dict["mars_input_path"]
        self.batch_processing = reader.control_dict['batch_processing']
        self.plot_swtich = reader.control_dict['plot_swtich']
        self.tracer_history = reader.control_dict['tracer_history']
        self.x_list = reader.x_list
        self.y_list = list(reader.y_list) + [-i for i in range(1, reader.core_properties_dict['discharge_mesh']+1)]
        self.discrete_mesh = reader.core_properties_dict['discharge_mesh']
        self.main_control = main_control
        # 初始化
        self.mesh_real_element_number = self.element_number / (len(self.x_list)*(len(self.y_list)-self.discrete_mesh))
        self._swtich_on(reader)
        self.target_time_str_list = []
        self.target_time_list = []
        self.energy_list = []
        
        if self.plot_swtich == 'on':
            
            self.release_dict_2d_plot = {}
            self.inventory_dict_2d_plot = {}
            self.isotope_dict_2d_plot = {}
            self.gamma_emission_2d_plot = {}
            self.gamma_emission_2d_plot_path = {}
            self.release_dict_2d_plot_path = {}
            self.inventory_dict_2d_plot_path = {}
            self.isotope_dict_2d_plot_path = {}
            self.release_dict_1d_plot = {}
            self.inventory_dict_1d_plot = {}
            self.isotope_dict_1d_plot = {}
            self.gamma_emission_1d_plot = {}
            self.release_dict_2d_max = {}
            self.inventory_dict_2d_max = {}
            self.isotope_dict_2d_max = {}
            self.release_dict_2d_min = {}
            self.inventory_dict_2d_min = {}
            self.isotope_dict_2d_min = {}
            self.gamma_emission_2d_min = {}
            self.gamma_emission_2d_max = {}
            self.decay_heat_2d_plot = []
            self.fission_power_2d_plot = []
            self.element_number_2d_plot = []
            self.decay_heat_2d_plot_path = []
            self.fission_power_2d_plot_path = []
            self.element_number_2d_plot_path = []
            self.decay_heat_1d_plot = []
            self.fission_power_1d_plot = []
            self.element_number_1d_plot = []
            self.decay_heat_2d_min = 1e30
            self.decay_heat_2d_max = -1.0
            self.fission_power_2d_min = 1e30
            self.fission_power_2d_max = -1.0
            self.element_number_2d_min = 1e30
            self.element_number_2d_max = -1.0
            for nuclide in self.output_nuclide:
                self.release_dict_2d_plot[nuclide] = []
                self.inventory_dict_2d_plot[nuclide] = []
                self.isotope_dict_2d_plot[nuclide] = []
                self.release_dict_1d_plot[nuclide] = []
                self.inventory_dict_1d_plot[nuclide] = []
                self.isotope_dict_1d_plot[nuclide] = []
                self.release_dict_2d_plot_path[nuclide] = []
                self.inventory_dict_2d_plot_path[nuclide] = []
                self.isotope_dict_2d_plot_path[nuclide] = []
                self.release_dict_2d_max[nuclide] = -1.0
                self.inventory_dict_2d_max[nuclide] = -1.0
                self.isotope_dict_2d_max[nuclide] = -1.0
                self.release_dict_2d_min[nuclide] = 1e30
                self.inventory_dict_2d_min[nuclide] = 1e30
                self.isotope_dict_2d_min[nuclide] = 1e30
                
    def make_dir(self,batch):
        '''
            创建子目录
        '''
        self.output_dir = os.path.join(self.output_dir_raw,f'batch{batch}')
        self.nuit_intermediate_dir = os.path.join(self.output_dir, 'NuitIntermediateFiles')
        self.mars_intermediate_dir = os.path.join(self.output_dir,'MarsIntermediateFiles')
        if self.main_control:
            if os.path.exists(self.output_dir):
                shutil.rmtree(self.output_dir)
            os.makedirs(self.output_dir)
            os.makedirs(self.nuit_intermediate_dir, exist_ok=True)
            if self.diffusion_swtich:
                os.makedirs(self.mars_intermediate_dir, exist_ok=True)
            self.finish_swtich = False
            self.finish_time = round_sig(sum(self.time_list))
            self.time_step = round_sig(self.time_step_list[0])
            self.time_span = round_sig(self.time_list[0])
            self.write_interval = round_sig(self.write_interval_list[0])
            self.write_time = self.write_interval
            self.time_tot = 0.0
            self.phase = 0
            self.step_tot = 0
            
    def _swtich_on(self,reader):
        '''
            读取数据的开关
        '''
        self.decay_heat_swtich = True if ('ReleaseRate' in reader.control_dict['output_type'] or 'FuelPerformance' in reader.control_dict['output_type'] or 'DecayHeat' in reader.control_dict['output_type']) else False
        self.gamma_spectrum_swtich = True if 'GammaSpectrum' in reader.control_dict['output_type'] else False
        self.isotope_swtich = True if 'Isotope' in reader.control_dict['output_type'] else False
        self.fission_power_swtich = True if ('ReleaseRate' in reader.control_dict['output_type'] or 'FuelPerformance' in reader.control_dict['output_type'] or 'FissionPower' in reader.control_dict['output_type'] or 'constpower' == reader.solver_dict['nuit_mode']) else False
        self.neutron_flux_swtich = True if ('ReleaseRate' in reader.control_dict['output_type'] or 'FuelPerformance' in reader.control_dict['output_type'] or 'constflux' == reader.solver_dict['nuit_mode']) else False
        self.read_temperature_swtich = True if ('ReleaseRate' in reader.control_dict['output_type'] or 'FuelPerformance' in reader.control_dict['output_type']) else False
        self.read_neutron_flux_swtich = True if 'constflux' == reader.solver_dict['nuit_mode'] else False
        self.read_power_swtich = True if 'constpower' == reader.solver_dict['nuit_mode'] else False
        self.diffusion_swtich = True if ('ReleaseRate' in reader.control_dict['output_type'] or 'FuelPerformance' in reader.control_dict['output_type']) else False
        
        self.delete_nuit_files = False if 'Nuit' in reader.control_dict['save_intermediate_files'] else True
        self.delete_mars_files = False if 'Mars' in reader.control_dict['save_intermediate_files'] else True
   
    def _count_mesh_cumulant_value(self,tracer_list:List[Tracer],property_name,nuclide=None,target_time=None,energy=None):
        '''
            统计网格累计值
        '''
        cumulant_matrix = np.zeros((len(self.x_list),len(self.y_list)),dtype=float)
        element_number_matrix = np.zeros((len(self.x_list),len(self.y_list)),dtype=int)
        element_number_real_matrix = np.zeros((len(self.x_list),len(self.y_list)),dtype=int)
        
        for tracer in tracer_list:
            time_list = tracer.time_sequence
            property_history = getattr(tracer,property_name)
            if nuclide is not None:
                property_history = property_history[nuclide]
            if energy is not None:
                property_history = property_history[energy]
            if len(property_history) > len(tracer.layer_location_history):
                property_history = np.asarray(property_history[1:])/2 + np.asarray(property_history[0:-1])/2
            if len(time_list) > len(property_history):
                time_list = np.array(tracer.time_sequence[0:-1])/2 + np.array(tracer.time_sequence[1:])/2
            if target_time is None:
                steady_idx = bisect.bisect_left(tracer.time_sequence,self.accident_begin)
                if tracer.time_sequence[-1] <= self.accident_begin:
                    property_history = property_history[0:steady_idx]
                    for idx,property_ in enumerate(property_history):
                        layer = tracer.layer_location_history[idx]
                        channel = tracer.channel_location_history[idx]
                        cumulant_matrix[channel][layer] += property_
                        element_number_matrix[channel][layer] += 1
                if tracer.time_sequence[-1] >= self.accident_begin and tracer.time_sequence[0] <= self.accident_begin:
                    layer = tracer.layer_location_history[steady_idx-1]
                    channel = tracer.channel_location_history[steady_idx-1]
                    element_number_real_matrix[channel][layer] += 1
            else:
                if target_time < tracer.time_sequence[0] or target_time > tracer.time_sequence[-1]:
                    continue
                idx = bisect.bisect_left(tracer.time_sequence,target_time)
                property_ = np.interp(target_time,time_list,property_history)
                layer = tracer.layer_location_history[idx-1]
                channel = tracer.channel_location_history[idx-1]
                cumulant_matrix[channel][layer] += property_
                element_number_matrix[channel][layer] += 1
                element_number_real_matrix = None    
        
        return cumulant_matrix,element_number_matrix,element_number_real_matrix

    def _output_tracer_history(self,tracer_list:List[Tracer]):
        '''
            输出Tracer的历史
        '''
        tracer_dir = os.path.join(self.output_dir,'TracerHistory')
        if not os.path.exists(tracer_dir):
            os.mkdir(tracer_dir)
        f_list = [open(os.path.join(tracer_dir,f'Tracer{tracer.idx}.txt'),'a') for tracer in tracer_list]
        f_discharge_list = [open(os.path.join(tracer_dir,f'Tracer{tracer.idx}Discharge.txt'),'a') for tracer in tracer_list]
        nuclides = self.output_nuclide
        properties = ['Time(s)','R(cm)','H(cm)']
        properties_discharge = ['Batch']
        if 'DecayHeat' in self.output_type:
            properties.append(f'DecayHeat(MW)')
        if 'FissionPower' in self.output_type:
            properties.append(f'FissionPower(MW)')
        if 'GammaSpectrum' in self.output_type:
            for energy in tracer_list[0].gamma_spectrum_history.keys():
                properties.append(f'Rate(n/s)-{float(energy):.1e}(Mev)')
                properties_discharge.append(f'DisRate(n/s){float(energy):.1e}(Mev)')
                properties_discharge.append(f'StoRate(n/s){float(energy):.1e}(Mev)')
        if 'ReleaseRate' in self.output_type:
            for nuclide in nuclides:
                properties.append(f'{nuclide}Release(Bq/s)')
        if 'Inventory' in self.output_type:
            for nuclide in nuclides:
                properties.append(f'{nuclide}Inventory(Bq)')
                properties_discharge.append(f'Dis{nuclide}Inv(Bq)')
                properties_discharge.append(f'Sto{nuclide}Inv(Bq)')
        if 'Isotope' in self.output_type:
            for nuclide in nuclides:
                properties.append(f'{nuclide}Isotope(atom)')
                properties_discharge.append(f'Dis{nuclide}Iso(atom)')
                properties_discharge.append(f'Sto{nuclide}Iso(atom)')
                
        for f,tracer in zip(f_list,tracer_list):
            f.write(formatting_line(properties))
            time_list = tracer.time_sequence
            time_list = np.array(tracer.time_sequence[0:-1])/2 + np.array(tracer.time_sequence[1:])/2
            for idx,time in enumerate(time_list):
                data = [time]
                data.append(self.x_list[tracer.channel_location_history[idx]])
                data.append(self.y_list[tracer.layer_location_history[idx]])
                if 'DecayHeat' in self.output_type:
                    data.append(tracer.decay_heat_history[idx])
                if 'FissionPower' in self.output_type:
                    data.append(tracer.power_history[idx])
                if 'GammaSpectrum' in self.output_type:
                    for energy,emission_history in tracer.gamma_spectrum_history.items():
                        data.append(emission_history[idx])
                if 'ReleaseRate' in self.output_type:
                    for nuclide in nuclides:
                        data.append(tracer.release_history_dict[nuclide][idx])
                if 'Inventory' in self.output_type:
                    for nuclide in nuclides:
                        data.append(tracer.inventory_history_dict[nuclide][idx]/2+tracer.inventory_history_dict[nuclide][idx+1]/2)
                if 'Isotope' in self.output_type:
                    for nuclide in nuclides:
                        data.append(tracer.isotope_history_dict[nuclide][idx]/2+tracer.isotope_history_dict[nuclide][idx+1]/2)
                
                f.write(formatting_line(data))
            f.close()
        
        for f_dis,tracer in zip(f_discharge_list,tracer_list):
            f_dis.write(formatting_line(properties_discharge))
            break_swtich = False
            for batch in range(999):
                data = [batch]
                if 'GammaSpectrum' in self.output_type:
                    for energy in tracer.discharge_gamma_spectrum_history.keys():
                        if batch < len(tracer.discharge_gamma_spectrum_history[energy]):
                            discharge = tracer.discharge_gamma_spectrum_history[energy][batch]
                            data.append(discharge)
                        else:
                            break_swtich = True
                            data.append('None')
                        if batch < len(tracer.storage_gamma_spectrum_history[energy]):
                            storage = tracer.storage_gamma_spectrum_history[energy][batch]
                            data.append(storage)
                        else:
                            break_swtich = True
                            data.append('None')
                if 'Inventory' in self.output_type:
                    for nuclide in tracer.discharge_inventory_history_dict.keys():
                        if batch < len(tracer.discharge_inventory_history_dict[nuclide]):
                            discharge = tracer.discharge_inventory_history_dict[nuclide][batch]
                            data.append(discharge)
                        else:
                            break_swtich = True
                            data.append('None')
                        if batch < len(tracer.storage_inventory_history_dict[nuclide]):
                            storage = tracer.storage_inventory_history_dict[nuclide][batch]
                            data.append(storage)
                        else:
                            break_swtich = True
                            data.append('None')
                if 'Isotope' in self.output_type:
                    for nuclide in tracer.discharge_isotope_history_dict.keys():
                        if batch < len(tracer.discharge_isotope_history_dict[nuclide]):
                            discharge = tracer.discharge_isotope_history_dict[nuclide][batch]
                            data.append(discharge)
                        else:
                            break_swtich = True
                            data.append('None')
                        if batch < len(tracer.storage_isotope_history_dict[nuclide]):
                            storage = tracer.storage_isotope_history_dict[nuclide][batch]
                            data.append(storage)
                        else:
                            break_swtich = True
                            data.append('None')
                f_dis.write(formatting_line(data))
                if break_swtich:
                    break
            f_dis.close()
                            
    def _cumulant_convert_to_mesh_average(self,cumulant_matrix,element_number_matrix):
        '''
            根据总元件数转换为网格平均值
        '''
        average_matrix = np.zeros_like(cumulant_matrix, dtype=float)
        np.divide(cumulant_matrix,element_number_matrix,out=average_matrix,where=(element_number_matrix != 0))
        average_matrix *= self.mesh_real_element_number
        
        return average_matrix
        
    def time_loop(self,tracer_list):
       
        if self.gamma_spectrum_swtich and len(self.energy_list) == 0 and self.plot_swtich == 'on':
            for energy in tracer_list[0].gamma_spectrum_history.keys():
                self.energy_list.append(energy)
                self.gamma_emission_2d_plot[energy] = []
                self.gamma_emission_2d_plot_path[energy] = []
                self.gamma_emission_1d_plot[energy] = []
                self.gamma_emission_2d_min[energy] = 1e30
                self.gamma_emission_2d_max[energy] = -1.0
        
        # 输出稳态结果
        self._output(tracer_list)
        
        while not self.finish_swtich:
            # 输出瞬态结果
            self._procced(tracer_list)

        # 输出Tracer历史
        if self.tracer_history == 'on':
            self._output_tracer_history(tracer_list)
        
    def _procced(self,tracer_list):
        '''
            推进时间步并记录数据
        '''
        if self.time_tot >= self.time_span:
            self.phase += 1
            if self.phase >= len(self.time_list):
                # 模拟结束
                self.finish_swtich = True
                return 
            else:
                # 进入下一阶段
                self.write_interval = self.write_interval_list[self.phase]
                self.time_span = self.time_span+self.time_list[self.phase]
                self.time_step = self.time_step_list[self.phase]
                if self.time_step >= self.time_list[self.phase]:
                    self.time_step = self.time_list[self.phase]/2
        if self.write_interval < 0:
            if self.step_tot % -self.write_interval == 0:
                # 记录当前信息
                if self.time_tot+self.time_step/2 >= self.accident_begin:
                    self._output(tracer_list,self.time_tot+self.time_step/2)
        elif self.write_interval > 0:
            if self.time_tot >= self.write_time:
                self.write_time += self.write_interval
                # 记录当前信息
                if self.time_tot+self.time_step/2 >= self.accident_begin:
                    self._output(tracer_list,self.time_tot+self.time_step/2)
        # 推进时间步
        if round_sig(self.time_tot + self.time_step) <= self.time_span:
            self.time_tot = self.time_tot + self.time_step
        else:
            self.time_step = self.time_span - self.time_tot
            self.time_tot = self.time_span
        self.step_tot = self.step_tot + 1
    
    def _output(self,tracer_list,target_time=None):
        '''
            输出
        '''
        if 'ReleaseRate' in self.output_type:
            self._output_release(tracer_list,target_time)
        if 'Inventory' in self.output_type:
            self._output_inventory(tracer_list,target_time)
        if 'Isotope' in self.output_type:
            self._output_isotope(tracer_list,target_time)
        if 'DecayHeat' in self.output_type:
            self._output_decay_heat(tracer_list,target_time)
        if 'FissionPower' in self.output_type:
            self._output_fission_power(tracer_list,target_time)
        if 'GammaSpectrum' in self.output_type:
            self._output_gamma_spectrum(tracer_list,target_time)
        
    def _output_inventory(self,tracer_list,target_time=None):
        '''
            输出盘存量
        '''
        for nuclide in self.output_nuclide:
            if target_time is None:
                cumulant_inventory_matrix,element_number_matrix,element_number_real_matrix = self._count_mesh_cumulant_value(tracer_list,'inventory_history_dict',nuclide=nuclide)
                inventory_matrix = self._cumulant_convert_to_mesh_average(cumulant_inventory_matrix,element_number_matrix)
                xlsx_path = os.path.join(self.output_dir,f'inventory_{nuclide}_steady.xlsx')
                self._save_two_matrices_to_xlsx(xlsx_path,inventory_matrix,A2=element_number_matrix,sheet1='Inventory',sheet2='ElementNumber',A3=element_number_real_matrix,sheet3='RealElementNumber',title='[Bq]')
            else:
                cumulant_inventory_matrix,element_number_matrix,_ = self._count_mesh_cumulant_value(tracer_list,'inventory_history_dict',nuclide=nuclide,target_time=target_time)
                inventory_matrix = self._cumulant_convert_to_mesh_average(cumulant_inventory_matrix,element_number_matrix)
                xlsx_path = os.path.join(self.output_dir,f'inventory_{nuclide}_transient_{target_time:.6e}.xlsx')
                self._save_two_matrices_to_xlsx(xlsx_path,inventory_matrix,A2=element_number_matrix,sheet1='Inventory',sheet2='ElementNumber',title='[Bq]')
    
    def _output_isotope(self,tracer_list,target_time=None):
        '''
            输出原子数量
        '''
        for nuclide in self.output_nuclide:
            if target_time is None:
                cumulant_isotope_matrix,element_number_matrix,element_number_real_matrix = self._count_mesh_cumulant_value(tracer_list,'isotope_history_dict',nuclide=nuclide)
                isotope_matrix = self._cumulant_convert_to_mesh_average(cumulant_isotope_matrix,element_number_matrix)
                xlsx_path = os.path.join(self.output_dir,f'isotope_{nuclide}_steady.xlsx')
                self._save_two_matrices_to_xlsx(xlsx_path,isotope_matrix,A2=element_number_matrix,sheet1='Isotope',sheet2='ElementNumber',A3=element_number_real_matrix,sheet3='RealElementNumber',title='[atom]')
            else:
                cumulant_isotope_matrix,element_number_matrix,_ = self._count_mesh_cumulant_value(tracer_list,'isotope_history_dict',nuclide=nuclide,target_time=target_time)
                isotope_matrix = self._cumulant_convert_to_mesh_average(cumulant_isotope_matrix,element_number_matrix)
                xlsx_path = os.path.join(self.output_dir,f'isotope_{nuclide}_transient_{target_time:.6e}.xlsx')
                self._save_two_matrices_to_xlsx(xlsx_path,isotope_matrix,A2=element_number_matrix,sheet1='Isotope',sheet2='ElementNumber',title='[atom]')
                   
    def _output_release(self,tracer_list,target_time=None):
        '''
            输出释放率
        '''
        for nuclide in self.output_nuclide:
            if target_time is None:
                cumulant_release_matrix,element_number_matrix,element_number_real_matrix = self._count_mesh_cumulant_value(tracer_list,'release_history_dict',nuclide=nuclide)
                release_matrix = self._cumulant_convert_to_mesh_average(cumulant_release_matrix,element_number_matrix)
                xlsx_path = os.path.join(self.output_dir,f'release_{nuclide}_steady.xlsx')
                self._save_two_matrices_to_xlsx(xlsx_path,release_matrix,A2=element_number_matrix,sheet1='Release',sheet2='ElementNumber',A3=element_number_real_matrix,sheet3='RealElementNumber',title='[Bq]/[s]')
            else:
                cumulant_release_matrix,element_number_matrix,_ = self._count_mesh_cumulant_value(tracer_list,'release_history_dict',nuclide=nuclide,target_time=target_time)
                release_matrix = self._cumulant_convert_to_mesh_average(cumulant_release_matrix,element_number_matrix)
                xlsx_path = os.path.join(self.output_dir,f'release_{nuclide}_transient_{target_time:.6e}.xlsx')
                self._save_two_matrices_to_xlsx(xlsx_path,release_matrix,A2=element_number_matrix,sheet1='Release',sheet2='ElementNumber',title='[Bq]/[s]')
        
    def _output_decay_heat(self,tracer_list,target_time=None):
        '''
            输出衰变热
        '''
        if target_time is None:
            cumulant_decay_heat_matrix,element_number_matrix,element_number_real_matrix = self._count_mesh_cumulant_value(tracer_list,'decay_heat_history')
            decay_heat_matrix = self._cumulant_convert_to_mesh_average(cumulant_decay_heat_matrix,element_number_matrix)
            xlsx_path = os.path.join(self.output_dir,f'decay_heat_steady.xlsx')
            self._save_two_matrices_to_xlsx(xlsx_path,decay_heat_matrix,A2=element_number_matrix,sheet1='DecayHeat',sheet2='ElementNumber',A3=element_number_real_matrix,sheet3='RealElementNumber',title='[MW]')
        else:
            cumulant_decay_heat_matrix,element_number_matrix,_ = self._count_mesh_cumulant_value(tracer_list,'decay_heat_history',target_time=target_time)
            decay_heat_matrix = self._cumulant_convert_to_mesh_average(cumulant_decay_heat_matrix,element_number_matrix)
            xlsx_path = os.path.join(self.output_dir,f'decay_heat_transient_{target_time:.6e}.xlsx')
            self._save_two_matrices_to_xlsx(xlsx_path,decay_heat_matrix,A2=element_number_matrix,sheet1='DecayHeat',sheet2='ElementNumber',title='[MW]')
        
    def _output_fission_power(self,tracer_list,target_time=None):
        '''
            输出裂变功率
        '''
        if target_time is None:
            cumulant_fission_power_matrix,element_number_matrix,element_number_real_matrix = self._count_mesh_cumulant_value(tracer_list,'power_history')
            fission_power_matrix = self._cumulant_convert_to_mesh_average(cumulant_fission_power_matrix,element_number_matrix)
            xlsx_path = os.path.join(self.output_dir,f'fission_power_steady.xlsx')
            self._save_two_matrices_to_xlsx(xlsx_path,fission_power_matrix,A2=element_number_matrix,sheet1='FissionPower',sheet2='ElementNumber',A3=element_number_real_matrix,sheet3='RealElementNumber',title='[MW]')
        else:
            cumulant_fission_power_matrix,element_number_matrix,_ = self._count_mesh_cumulant_value(tracer_list,'power_history',target_time=target_time)
            fission_power_matrix = self._cumulant_convert_to_mesh_average(cumulant_fission_power_matrix,element_number_matrix)
            xlsx_path = os.path.join(self.output_dir,f'fission_power_transient_{target_time:.6e}.xlsx')
            self._save_two_matrices_to_xlsx(xlsx_path,fission_power_matrix,A2=element_number_matrix,sheet1='FissionPower',sheet2='ElementNumber',title='[MW]')
    
    def _output_gamma_spectrum(self,tracer_list,target_time=None):
        '''
            输出光子释放率
        '''
        for energy in tracer_list[0].gamma_spectrum_history.keys():
            if target_time is None:
                cumulant_emission_rate_matrix,element_number_matrix,element_number_real_matrix = self._count_mesh_cumulant_value(tracer_list,'gamma_spectrum_history',energy=energy)
                emission_rate_matrix = self._cumulant_convert_to_mesh_average(cumulant_emission_rate_matrix,element_number_matrix)
                xlsx_path = os.path.join(self.output_dir,f'gamma_emission_{energy}Mev_steady.xlsx')
                self._save_two_matrices_to_xlsx(xlsx_path,emission_rate_matrix,A2=element_number_matrix,sheet1='GammaEmission',sheet2='ElementNumber',A3=element_number_real_matrix,sheet3='RealElementNumber',title='[n]/[s]')
            else:
                cumulant_emission_rate_matrix,element_number_matrix,element_number_real_matrix = self._count_mesh_cumulant_value(tracer_list,'gamma_spectrum_history',energy=energy,target_time=target_time)
                emission_rate_matrix = self._cumulant_convert_to_mesh_average(cumulant_emission_rate_matrix,element_number_matrix)
                xlsx_path = os.path.join(self.output_dir,f'gamma_emission_{energy}Mev_transient_{target_time:.6e}.xlsx')
                self._save_two_matrices_to_xlsx(xlsx_path,emission_rate_matrix,A2=element_number_matrix,sheet1='GammaEmission',sheet2='ElementNumber',A3=element_number_real_matrix,sheet3='RealElementNumber',title='[n]/[s]')
                
    def _write_matrix_sheet(self,ws, A, x_list, y_list, title="R/Z"):
        '''
            输出矩阵到xlsx
        '''
        A = np.asarray(A)
        nx, ny = A.shape
        # 左上角
        ws.cell(row=1, column=1, value=title)
        # 顶部 x 坐标（列标题）
        for jx, x in enumerate(x_list, start=2):
            ws.cell(row=1, column=jx, value=x)
        # 左侧 y 坐标（行标题）
        for iy, y in enumerate(y_list, start=2):
            ws.cell(row=iy, column=1, value=y)
        # 写数据：行对应 y_index，列对应 x_index
        for iy in range(ny):
            for jx in range(nx):
                ws.cell(row=2 + iy, column=2 + jx, value=float(A[jx, iy]))      
        
    def _save_two_matrices_to_xlsx(self,xlsx_path,A1,sheet1="Matrix1",A2=None,sheet2="Matrix2",A3=None,sheet3="Matrix3",title="R/Z"):
        '''
            输出两个矩阵到xlsx
        '''
        x1 = self.x_list
        y1 = self.y_list
        wb = Workbook()
        ws1 = wb.active
        ws1.title = sheet1
        self._write_matrix_sheet(ws1, A1, x1, y1, title=title)

        if A2 is not None:
            ws2 = wb.create_sheet(title=sheet2)
            self._write_matrix_sheet(ws2, A2, x1, y1, title=title)
        if A3 is not None:
            ws3 = wb.create_sheet(title=sheet3)
            self._write_matrix_sheet(ws3, A3, x1, y1, title=title)

        wb.save(xlsx_path)
    
    def split_to_batch(self,tracer_list):
        '''
            将列表等量的差分为子列表
        '''
        n = self.batch_processing
        L = len(tracer_list)
        n = min(n, L)
        base, extra = divmod(L, n)
        self.tracer_batch_list = []
        start = 0
        for i in range(n):
            size = base + (1 if i < extra else 0)
            self.tracer_batch_list.append(tracer_list[start:start + size])
            start += size
    
    def _extract_and_merge_matrix(self,xlsx_path_list,sheet_name):
        '''
            提取并融合每个批次的矩阵
        '''
        merge_matrix = np.zeros((len(self.y_list),len(self.x_list)))
        merge_element_number_matrix = np.zeros((len(self.y_list),len(self.x_list)))
        merge_element_number_real_matrix = np.zeros((len(self.y_list),len(self.x_list)))
        for xlsx_path in xlsx_path_list:
            df = pd.read_excel(xlsx_path,header=None,sheet_name=sheet_name)
            matrix = df.iloc[1:,1:].to_numpy()
            df = pd.read_excel(xlsx_path,header=None,sheet_name='ElementNumber')
            if 'steady' in xlsx_path:
                df2 = pd.read_excel(xlsx_path,header=None,sheet_name='RealElementNumber')
                element_number_real_matrix = df2.iloc[1:,1:].to_numpy()
                merge_element_number_real_matrix += element_number_real_matrix
            element_number_matrix = df.iloc[1:,1:].to_numpy() 
            with np.errstate(divide='ignore', invalid='ignore'):
                matrix = matrix / self.mesh_real_element_number * element_number_matrix
            merge_matrix += matrix
            merge_element_number_matrix += element_number_matrix
        with np.errstate(divide='ignore', invalid='ignore'):
            merge_matrix = merge_matrix / merge_element_number_matrix * self.mesh_real_element_number
        total_value = np.nansum(merge_matrix[:-self.discrete_mesh, :]) # 卸出堆芯的部分不参与统计
        return merge_matrix,merge_element_number_matrix,total_value,merge_element_number_real_matrix
     
    def _make_xlsx_path_list(self,batch_paths,xlsx_name):
        '''
            制作xlsx的索引列表
        '''
        
        return [os.path.join(dir_p,xlsx_name) for dir_p in batch_paths]
    
    def _plot_append(self,merge_matrix,xlsx_name_raw,nuclide,total_value,target_time,energy=None):
        '''
            添加数据到绘画列表
        '''
        if 'inventory' in xlsx_name_raw:
            inventory_matrix = merge_matrix
            self.inventory_dict_2d_plot[nuclide].append(inventory_matrix)
            max_val = -1 if np.isnan(inventory_matrix).all() else np.nanmax(inventory_matrix)
            min_val = 1e30 if np.isnan(inventory_matrix).all() else np.nanmin(inventory_matrix)
            if max_val > self.inventory_dict_2d_max[nuclide]:
                self.inventory_dict_2d_max[nuclide] = max_val
            if min_val < self.inventory_dict_2d_min[nuclide]:
                self.inventory_dict_2d_min[nuclide] = min_val
            self.inventory_dict_1d_plot[nuclide].append(total_value)
            if target_time == 'steady':
                self.inventory_dict_1d_plot[nuclide].append(total_value)
        if 'isotope' in xlsx_name_raw:
            isotope_matrix = merge_matrix
            self.isotope_dict_2d_plot[nuclide].append(isotope_matrix)
            max_val = -1 if np.isnan(isotope_matrix).all() else np.nanmax(isotope_matrix)
            min_val = 1e30 if np.isnan(isotope_matrix).all() else np.nanmin(isotope_matrix)
            if max_val > self.isotope_dict_2d_max[nuclide]:
                self.isotope_dict_2d_max[nuclide] = max_val
            if min_val < self.isotope_dict_2d_min[nuclide]:
                self.isotope_dict_2d_min[nuclide] = min_val
            self.isotope_dict_1d_plot[nuclide].append(total_value)
            if target_time == 'steady':
                self.isotope_dict_1d_plot[nuclide].append(total_value)
        if 'release' in xlsx_name_raw:
            release_matrix = merge_matrix
            self.release_dict_2d_plot[nuclide].append(release_matrix)
            max_val = -1 if np.isnan(release_matrix).all() else np.nanmax(release_matrix)
            min_val = 1e30 if np.isnan(release_matrix).all() else np.nanmin(release_matrix)
            if max_val > self.release_dict_2d_max[nuclide]:
                self.release_dict_2d_max[nuclide] = max_val
            if min_val < self.release_dict_2d_min[nuclide]:
                self.release_dict_2d_min[nuclide] = min_val
            self.release_dict_1d_plot[nuclide].append(total_value)
            if target_time == 'steady':
                self.release_dict_1d_plot[nuclide].append(total_value)
        if 'decay_heat' in xlsx_name_raw:
            decay_heat_matrix = merge_matrix
            self.decay_heat_2d_plot.append(decay_heat_matrix)
            max_val = -1 if np.isnan(decay_heat_matrix).all() else np.nanmax(decay_heat_matrix)
            min_val = 1e30 if np.isnan(decay_heat_matrix).all() else np.nanmin(decay_heat_matrix)
            if max_val > self.decay_heat_2d_max:
                self.decay_heat_2d_max = max_val
            if min_val < self.decay_heat_2d_min:
                self.decay_heat_2d_min = min_val
            self.decay_heat_1d_plot.append(total_value)
            if target_time == 'steady':
                self.decay_heat_1d_plot.append(total_value)
        if 'fission_power' in xlsx_name_raw:
            fission_power_matrix = merge_matrix
            self.fission_power_2d_plot.append(fission_power_matrix)
            max_val = -1 if np.isnan(fission_power_matrix).all() else np.nanmax(fission_power_matrix)
            min_val = 1e30 if np.isnan(fission_power_matrix).all() else np.nanmin(fission_power_matrix)
            if max_val > self.fission_power_2d_max:
                self.fission_power_2d_max = max_val
            if min_val < self.fission_power_2d_min:
                self.fission_power_2d_min = min_val
            self.fission_power_1d_plot.append(total_value)
            if target_time == 'steady':
                self.fission_power_1d_plot.append(total_value)
        if 'gamma_emission' in xlsx_name_raw:
            gamma_emission_matrix = merge_matrix
            self.gamma_emission_2d_plot[energy].append(gamma_emission_matrix)
            max_val = -1 if np.isnan(gamma_emission_matrix).all() else np.nanmax(gamma_emission_matrix)
            min_val = 1e30 if np.isnan(gamma_emission_matrix).all() else np.nanmin(gamma_emission_matrix)
            if max_val > self.gamma_emission_2d_max[energy]:
                self.gamma_emission_2d_max[energy] = max_val
            if min_val < self.gamma_emission_2d_min[energy]:
                self.gamma_emission_2d_min[energy] = min_val
            self.gamma_emission_1d_plot[energy].append(total_value)
            if target_time == 'steady':
                self.gamma_emission_1d_plot[energy].append(total_value)
        if 'element_number' in xlsx_name_raw and 'real' not in xlsx_name_raw:
            element_number_matrix = merge_matrix
            self.element_number_2d_plot.append(element_number_matrix)
            max_val = -1 if np.isnan(element_number_matrix).all() else np.nanmax(element_number_matrix)
            min_val = 1e30 if np.isnan(element_number_matrix).all() else np.nanmin(element_number_matrix)
            if max_val > self.element_number_2d_max:
                self.element_number_2d_max = max_val
            if min_val < self.element_number_2d_min:
                self.element_number_2d_min = min_val
            self.element_number_1d_plot.append(total_value)
            if target_time == 'steady':
                self.element_number_1d_plot.append(total_value)
                  
    def _write_merge_matrix(self,xlsx_name_raw,sheet_name,batch_paths,header,nuclide=None,energy=None):
        '''
            处理一个输出类型
        '''
        f = open(os.path.join(self.output_dir_raw,f'{xlsx_name_raw}.txt'),'a',buffering=1)
        f.write(header)
        xlsx_name_steady = xlsx_name_raw + '_steady.xlsx'
        xlsx_list_steady = self._make_xlsx_path_list(batch_paths,xlsx_name_steady)
        merge_matrix,merge_element_number_matrix,total_value,merge_element_number_real_matrix = self._extract_and_merge_matrix(xlsx_list_steady,sheet_name)
        if self.plot_swtich == 'on':
            self._plot_append(merge_matrix,xlsx_name_raw,nuclide,total_value,'steady',energy)
            if self.element_number_append:
                self._plot_append(merge_element_number_real_matrix,'element_number',nuclide,np.sum(merge_element_number_real_matrix),'steady',energy)
        self._save_two_matrices_to_xlsx(os.path.join(self.output_dir_raw,xlsx_name_steady),merge_matrix.T,sheet_name,merge_element_number_matrix.T,'ElementNumber',A3=merge_element_number_real_matrix.T,sheet3='ElementNumberReal')
        f.write(formatting_line(['steady',total_value]))
        for target_time in self.target_time_str_list:
            if 'steady' in target_time:
                continue
            xlsx_name_transient = xlsx_name_raw + f'_transient_{target_time}.xlsx'
            xlsx_list_transient = self._make_xlsx_path_list(batch_paths,xlsx_name_transient)
            merge_matrix,merge_element_number_matrix,total_value,_ = self._extract_and_merge_matrix(xlsx_list_transient,sheet_name)
            self._save_two_matrices_to_xlsx(os.path.join(self.output_dir_raw,xlsx_name_transient),merge_matrix.T,sheet_name,merge_element_number_matrix.T,'ElementNumber')
            f.write(formatting_line(['time',target_time,total_value]))
            if self.plot_swtich == 'on':
                self._plot_append(merge_matrix,xlsx_name_raw,nuclide,total_value,target_time,energy)  
                if self.element_number_append:
                    self._plot_append(merge_element_number_matrix,'element_number',nuclide,np.sum(merge_element_number_matrix),target_time,energy)
        self.element_number_append = False
        f.close()
        
        return merge_element_number_matrix
    
    def _time_loop_for_merge(self):
        '''
            为了融合批次再进行一次循环获得时间列表
        '''
        self.finish_swtich = False
        self.finish_time = round_sig(sum(self.time_list))
        self.time_step = round_sig(self.time_step_list[0])
        self.time_span = round_sig(self.time_list[0])
        self.write_interval = round_sig(self.write_interval_list[0])
        self.write_time = self.write_interval
        self.time_tot = 0.0
        self.phase = 0
        self.step_tot = 0
        self.target_time_str_list.append('steady')
        while not self.finish_swtich:
            if self.time_tot >= self.time_span:
                self.phase += 1
                if self.phase >= len(self.time_list):
                    # 模拟结束
                    self.finish_swtich = True
                    return
                else:
                    # 进入下一阶段
                    self.write_interval = self.write_interval_list[self.phase]
                    self.time_span = self.time_span+self.time_list[self.phase]
                    self.time_step = self.time_step_list[self.phase]
                    if self.time_step >= self.time_list[self.phase]:
                        self.time_step = self.time_list[self.phase]/2
            if self.write_interval < 0:
                if self.step_tot % -self.write_interval == 0:
                    # 记录当前信息
                    if self.time_tot+self.time_step/2 >= self.accident_begin:
                        self.target_time_str_list.append(f'{self.time_tot+self.time_step/2:.6e}')
            elif self.write_interval > 0:
                if self.time_tot >= self.write_time:
                    self.write_time += self.write_interval
                    # 记录当前信息
                    if self.time_tot+self.time_step/2 >= self.accident_begin:
                        self.target_time_str_list.append(f'{self.time_tot+self.time_step/2:.6e}')
            # 推进时间步
            if round_sig(self.time_tot + self.time_step) <= self.time_span:
                self.time_tot = self.time_tot + self.time_step
            else:
                self.time_step = self.time_span - self.time_tot
                self.time_tot = self.time_span
            self.step_tot = self.step_tot + 1
    
    def merge_batch(self,tracer):
        '''
            融合各个批次
        '''
        root = Path(self.output_dir_raw)
        batch_dirs = sorted(p for p in root.glob("batch*") if p.is_dir())
        batch_paths = [str(p.resolve()) for p in batch_dirs]
        self._time_loop_for_merge()
        self.element_number_append = True
        for target_time in self.target_time_str_list:
            if target_time == 'steady':
                self.target_time_list.append(0.0)
                time = self.finish_time if self.accident_begin > self.finish_time else self.accident_begin
                self.target_time_list.append(float(time))
            else:
                self.target_time_list.append(float(target_time))
        if 'Inventory' in self.output_type: 
            for nuclide in self.output_nuclide:
                self._write_merge_matrix(f'inventory_{nuclide}','Inventory',batch_paths,f'Full core {nuclide} inventory ([Bq])\n',nuclide)
        if 'Isotope' in self.output_type: 
            for nuclide in self.output_nuclide:
                self._write_merge_matrix(f'isotope_{nuclide}','Isotope',batch_paths,f'Full core {nuclide} isotope ([atom])\n',nuclide)
        if 'ReleaseRate' in self.output_type: 
            for nuclide in self.output_nuclide:
                self._write_merge_matrix(f'release_{nuclide}','Release',batch_paths,f'Full core {nuclide} release rate ([Bq]/[s])\n',nuclide)
        if 'DecayHeat' in self.output_type: 
            self._write_merge_matrix(f'decay_heat','DecayHeat',batch_paths,f'Full core decay heat ([MW])\n')
        if 'FissionPower' in self.output_type: 
            self._write_merge_matrix(f'fission_power','FissionPower',batch_paths,f'Full core fission power ([MW])\n')
        if 'GammaSpectrum' in self.output_type: 
            self.energy_list = list(tracer.gamma_spectrum_history.keys())
            for energy in self.energy_list:
                self._write_merge_matrix(f'gamma_emission_{energy}Mev','GammaEmission',batch_paths,f'Full core gamma emission ([n]/[s])\n',energy=energy)
        # 绘图
        if self.plot_swtich == 'on':
            self.Plot.init_plot(self.output_dir_raw,self.accident_begin,self.x_list,self.y_list)
            self._plot_output()
        # 删除批次目录
        #for dir_p in batch_paths:
        #    shutil.rmtree(dir_p)  
            
    def _plot_output(self):
        '''
            对输出结果绘图
        '''
        # 2d图
        for idx,targe_time in enumerate(self.target_time_str_list):
            if 'ReleaseRate' in self.output_type:
                for nuclide in self.output_nuclide:    
                    self.Plot.plot_2d(self.release_dict_2d_plot[nuclide][idx],f'{nuclide}ReleaseRate','([Bq]/[s])',self.release_dict_2d_max[nuclide],self.release_dict_2d_min[nuclide],targe_time,f'release_rate_{nuclide}')
                    self.release_dict_2d_plot_path[nuclide].append(os.path.join(self.Plot.output_dir,f'release_rate_{nuclide}_{targe_time}'+'.png'))
            if 'Inventory' in self.output_type:
                for nuclide in self.output_nuclide:
                    self.Plot.plot_2d(self.inventory_dict_2d_plot[nuclide][idx],f'{nuclide}Inventory','([Bq])',self.inventory_dict_2d_max[nuclide],self.inventory_dict_2d_min[nuclide],targe_time,f'inventory_{nuclide}')
                    self.inventory_dict_2d_plot_path[nuclide].append(os.path.join(self.Plot.output_dir,f'inventory_{nuclide}_{targe_time}'+'.png'))
            if 'Isotope' in self.output_type:
                for nuclide in self.output_nuclide:
                    self.Plot.plot_2d(self.isotope_dict_2d_plot[nuclide][idx],f'{nuclide}Isotope','([atom])',self.isotope_dict_2d_max[nuclide],self.isotope_dict_2d_min[nuclide],targe_time,f'isotope_{nuclide}')
                    self.isotope_dict_2d_plot_path[nuclide].append(os.path.join(self.Plot.output_dir,f'isotope_{nuclide}_{targe_time}'+'.png'))
            if 'DecayHeat' in self.output_type:
                self.Plot.plot_2d(self.decay_heat_2d_plot[idx],f'DecayHeat','([MW])',self.decay_heat_2d_max,self.decay_heat_2d_min,targe_time,f'decay_heat')
                self.decay_heat_2d_plot_path.append(os.path.join(self.Plot.output_dir,f'decay_heat_{targe_time}'+'.png'))
            if 'FissionPower' in self.output_type:
                self.Plot.plot_2d(self.fission_power_2d_plot[idx],f'FissionPower','([MW])',self.fission_power_2d_max,self.fission_power_2d_min,targe_time,f'fission_power')
                self.fission_power_2d_plot_path.append(os.path.join(self.Plot.output_dir,f'fission_power_{targe_time}'+'.png'))
            if 'GammaSpectrum' in self.output_type:
                for energy in self.energy_list:
                    self.Plot.plot_2d(self.gamma_emission_2d_plot[energy][idx],f'{energy}MevEmissionRate','([n]/[s])',self.gamma_emission_2d_max[energy],self.gamma_emission_2d_min[energy],targe_time,f'Energy_{energy}')
                    self.gamma_emission_2d_plot_path[energy].append(os.path.join(self.Plot.output_dir,f'energy_{energy}_{targe_time}'+'.png'))
            self.Plot.plot_2d(self.element_number_2d_plot[idx],f'ElementNumber','',self.element_number_2d_max,self.element_number_2d_min,targe_time,f'element_number')
            self.element_number_2d_plot_path.append(os.path.join(self.Plot.output_dir,f'element_number_{targe_time}'+'.png'))
        # 1d图
        if 'ReleaseRate' in self.output_type:
            for nuclide in self.output_nuclide:
                self.Plot.plot_line(f'{nuclide}ReleaseRate','([Bq]/[s])',self.release_dict_1d_plot[nuclide],self.target_time_list,f'{nuclide}_release_rate') 
        if 'Inventory' in self.output_type:
            for nuclide in self.output_nuclide:
                self.Plot.plot_line(f'{nuclide}Inventory','([Bq])',self.inventory_dict_1d_plot[nuclide],self.target_time_list,f'{nuclide}_inventory')
        if 'Isotope' in self.output_type:
            for nuclide in self.output_nuclide:
                self.Plot.plot_line(f'{nuclide}Isotope','([atom])',self.isotope_dict_1d_plot[nuclide],self.target_time_list,f'{nuclide}_isotope')    
        if 'DecayHeat' in self.output_type:
            self.Plot.plot_line(f'DecayHeat','([MW])',self.decay_heat_1d_plot,self.target_time_list,f'decay_heat')
        if 'FissionPower' in self.output_type:
            self.Plot.plot_line(f'FissionPower','([MW])',self.fission_power_1d_plot,self.target_time_list,f'fission_power')
        if 'GammaSpectrum' in self.output_type:
            for energy in self.energy_list:
                self.Plot.plot_line(f'GammaSpectrum','([n]/[s])',self.gamma_emission_1d_plot[energy],self.target_time_list,f'gamma_emission_{energy}Mev')
        self.Plot.plot_line(f'ElementNumber','',self.element_number_1d_plot,self.target_time_list,f'element_number')
        if len(self.target_time_str_list) > 1:
            # 瞬态2d图转gif
            if 'ReleaseRate' in self.output_type:
                for nuclide in self.output_nuclide:
                    self.Plot._png_to_gif(self.release_dict_2d_plot_path[nuclide],f'{nuclide}_release_rate')
                    for file_path in self.release_dict_2d_plot_path[nuclide]:
                        if 'steady' in file_path:
                            continue
                        os.remove(file_path)
            if 'Inventory' in self.output_type:
                for nuclide in self.output_nuclide:
                    self.Plot._png_to_gif(self.inventory_dict_2d_plot_path[nuclide],f'{nuclide}_inventory')
                    for file_path in self.inventory_dict_2d_plot_path[nuclide]:
                        if 'steady' in file_path:
                            continue
                        os.remove(file_path)
            if 'Isotope' in self.output_type:
                for nuclide in self.output_nuclide:
                    self.Plot._png_to_gif(self.isotope_dict_2d_plot_path[nuclide],f'{nuclide}_isotope')
                    for file_path in self.isotope_dict_2d_plot_path[nuclide]:
                        if 'steady' in file_path:
                            continue
                        os.remove(file_path)
            if 'DecayHeat' in self.output_type:
                self.Plot._png_to_gif(self.decay_heat_2d_plot_path,f'decay_heat')
                for file_path in self.decay_heat_2d_plot_path:
                    if 'steady' in file_path:
                        continue
                    os.remove(file_path)
            if 'FissionPower' in self.output_type:
                self.Plot._png_to_gif(self.fission_power_2d_plot_path,f'fission_power')
                for file_path in self.fission_power_2d_plot_path:
                    if 'steady' in file_path:
                        continue
                    os.remove(file_path)
            if 'GammaSpectrum' in self.output_type:
                for energy in self.energy_list:
                    self.Plot._png_to_gif(self.gamma_emission_2d_plot_path[energy],f'gamma_emission_{energy}Mev')
                    for file_path in self.gamma_emission_2d_plot_path[energy]:
                        if 'steady' in file_path:
                            continue
                        os.remove(file_path)
            self.Plot._png_to_gif(self.element_number_2d_plot_path,f'element_number')
        for file_path in self.element_number_2d_plot_path:
            if 'steady' in file_path:
                continue
            os.remove(file_path)